"""ppt_builder.py — AI 마크다운 보고서 → PPT 변환기"""
import io
import re
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

NAVY  = RGBColor(0x1E, 0x3A, 0x8A)
BLUE  = RGBColor(0x25, 0x63, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK  = RGBColor(0x1E, 0x29, 0x3B)
GRAY  = RGBColor(0x64, 0x74, 0x8B)
LGRAY = RGBColor(0xF1, 0xF5, 0xF9)

SECTION_COLORS = [
    RGBColor(0x25, 0x63, 0xEB),  # 파랑
    RGBColor(0xD9, 0x77, 0x06),  # 앰버
    RGBColor(0x16, 0xA3, 0x4A),  # 초록
]

SW = Inches(13.33)
SH = Inches(7.5)


def I(inch): return Inches(inch)
def P(pt):   return Pt(pt)


def _shape(slide, x, y, w, h, fill=None, line_color=None):
    s = slide.shapes.add_shape(1, I(x), I(y), I(w), I(h))
    s.line.fill.background()
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line_color:
        s.line.color.rgb = line_color
    else:
        s.line.fill.background()
    return s


def _tb(slide, x, y, w, h, text, size=11, bold=False, color=None,
        align=PP_ALIGN.LEFT):
    txb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf  = txb.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text  = text
    run.font.size = P(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    return txb


def _parse_sections(markdown: str) -> list:
    parts = re.split(r'(?=^## \d+\.)', markdown, flags=re.MULTILINE)
    sections = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        m = re.match(r'^## (\d+)\.\s*(.+)', part)
        if m:
            title   = f"{m.group(1)}. {m.group(2).strip()}"
            body    = part[m.end():].strip()
            bullets = _to_bullets(body)
            sections.append((title, bullets))
    return sections


def _clean(text: str) -> str:
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+)\*', r'\1', text)
    return text.strip()


def _to_bullets(text: str) -> list:
    lines   = text.split('\n')
    bullets = []
    para    = []

    def flush():
        if para:
            t = _clean(' '.join(para))
            if t:
                bullets.append(('p', t))
            para.clear()

    for line in lines:
        t = line.strip()
        if not t:
            flush()
            continue
        # 소제목: **텍스트**
        if re.match(r'^\*\*[^*]+\*\*$', t):
            flush()
            bullets.append(('h', _clean(t)))
        # 글머리
        elif re.match(r'^[*\-]\s+', t):
            flush()
            bullets.append(('b', _clean(re.sub(r'^[*\-]\s+', '', t))))
        # 번호 항목
        elif re.match(r'^\d+\.\s+', t):
            flush()
            bullets.append(('n', _clean(re.sub(r'^\d+\.\s+', '', t))))
        else:
            para.append(_clean(t))
    flush()
    return bullets


def _title_slide(prs, today: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _shape(slide, 0, 0, 13.33, 7.5, fill=NAVY)
    _shape(slide, 0, 3.4, 13.33, 0.06, fill=BLUE)

    _tb(slide, 1.2, 1.8, 10.9, 0.9,
        'Field Health 불량 예측 분석 보고서',
        size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _tb(slide, 1.2, 2.8, 10.9, 0.4,
        'AI 기반 공정 개선 방향 보고서',
        size=15, color=RGBColor(0x93, 0xC5, 0xFD), align=PP_ALIGN.CENTER)
    _tb(slide, 1.2, 3.65, 10.9, 0.3,
        today, size=12,
        color=RGBColor(0xCB, 0xD5, 0xE1), align=PP_ALIGN.CENTER)

    _shape(slide, 0, 7.1, 13.33, 0.4, fill=RGBColor(0x0F, 0x17, 0x2A))
    _tb(slide, 0.4, 7.15, 7, 0.25,
        'SK hynix  |  DRAM Wafer Test Field Health Prediction',
        size=8, color=RGBColor(0x94, 0xA3, 0xB8))
    _tb(slide, 7, 7.15, 6, 0.25,
        'We Do Technology  |  Confidential',
        size=8, bold=True, color=RGBColor(0x93, 0xC5, 0xFD),
        align=PP_ALIGN.RIGHT)


def _content_slide(prs, section_num: int, title: str,
                   bullets: list, accent: RGBColor, today: str, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _shape(slide, 0, 0, 13.33, 7.5, fill=WHITE)
    _shape(slide, 0, 0, 13.33, 0.56, fill=accent)
    _tb(slide, 0.3, 0.1, 11.8, 0.38, title,
        size=17, bold=True, color=WHITE)
    _tb(slide, 12.0, 0.1, 1.1, 0.38,
        f'{section_num} / {total}', size=9,
        color=WHITE, align=PP_ALIGN.RIGHT)

    y     = 0.72
    MAX_Y = 6.90
    LM    = 0.45
    CW    = 12.43

    for kind, text in bullets:
        if y >= MAX_Y:
            break

        if kind == 'h':
            _shape(slide, LM, y, CW, 0.28, fill=LGRAY)
            _tb(slide, LM + 0.1, y + 0.02, CW - 0.2, 0.26,
                text, size=10, bold=True, color=DARK)
            y += 0.33
        elif kind == 'b':
            _tb(slide, LM, y, CW, 0.32,
                f'• {text}', size=10, color=DARK)
            y += 0.32
        elif kind == 'n':
            _tb(slide, LM, y, CW, 0.32,
                f'  → {text}', size=10, color=DARK)
            y += 0.32
        else:
            _tb(slide, LM, y, CW, 0.32, text, size=10,
                color=RGBColor(0x47, 0x55, 0x69))
            y += 0.32

    _shape(slide, 0, 7.1, 13.33, 0.02,
           fill=RGBColor(0xCB, 0xD5, 0xE1))
    _tb(slide, 0.3, 7.13, 8, 0.22,
        f'{today}  Field Health Prediction Model  Two-Stage',
        size=7, color=GRAY)
    _tb(slide, 8, 7.13, 5, 0.22,
        'SK hynix  |  대외비', size=7, bold=True,
        color=NAVY, align=PP_ALIGN.RIGHT)


def build_text_ppt(markdown: str) -> bytes:
    """AI 마크다운 보고서를 PPT로 변환하여 bytes 반환"""
    prs = Presentation()
    prs.slide_width  = SW
    prs.slide_height = SH

    today    = datetime.now().strftime('%Y년 %m월 %d일')
    sections = _parse_sections(markdown)

    _title_slide(prs, today)

    if not sections:
        # 섹션 구분이 없으면 단일 슬라이드로
        bullets = _to_bullets(markdown)
        _content_slide(prs, 1, '보고서 내용', bullets,
                       SECTION_COLORS[0], today, 1)
    else:
        total = len(sections)
        for i, (title, bullets) in enumerate(sections):
            color = SECTION_COLORS[i % len(SECTION_COLORS)]
            _content_slide(prs, i + 1, title, bullets, color, today, total)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()
