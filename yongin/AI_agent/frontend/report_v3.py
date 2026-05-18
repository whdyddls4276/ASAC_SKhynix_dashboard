#!/usr/bin/env python3
"""report_v3.py — Field Health 불량 예측 분석 보고서 v3 PPT 생성기 (완전 네이티브 차트)

차트 5종 모두 PowerPoint 내장 차트 / 도형으로 생성됩니다.
  - 트렌드 라인차트    → LINE (우클릭 → 데이터 편집)
  - SHAP 막대         → BAR_CLUSTERED (우클릭 → 데이터 편집)
  - 피처 임포턴스     → BAR_CLUSTERED (우클릭 → 데이터 편집)
  - 박스플롯          → pptx 도형 (클릭으로 색상·위치·크기 수정)
  - 히스토그램        → COLUMN_CLUSTERED (우클릭 → 데이터 편집)

사용법:
    python report_v3.py
    python report_v3.py --data report_data.json --out report_v3.pptx
"""

import json, os, sys, argparse
from datetime import date, timedelta

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from lxml import etree

# ── 색상 팔레트 ────────────────────────────────────────────
RED    = RGBColor(0xEF, 0x44, 0x44)
GREEN  = RGBColor(0x22, 0xC5, 0x5E)
BLUE   = RGBColor(0x3B, 0x82, 0xF6)
AMBER  = RGBColor(0xF5, 0x9E, 0x0B)
GRAY   = RGBColor(0x94, 0xA3, 0xB8)
NAVY   = RGBColor(0x1E, 0x3A, 0x8A)
DARK   = RGBColor(0x0F, 0x17, 0x2A)
SLATE  = RGBColor(0x47, 0x55, 0x69)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LTBLUE = RGBColor(0xEF, 0xF6, 0xFF)
LTGRAY = RGBColor(0xF8, 0xFA, 0xFC)
BORDRC = RGBColor(0x94, 0xA3, 0xB8)
HDRBG  = RGBColor(0xD1, 0xD9, 0xE8)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
INDIGO = RGBColor(0x63, 0x66, 0xF1)

SW = Inches(13.33)
SH = Inches(7.5)

# ── 기본 유틸 ──────────────────────────────────────────────
def I(inch): return Inches(inch)
def P(pt):   return Pt(pt)

def day_to_date(d, last_day, y=2026, m=6, day=11):
    ref = date(y, m, day)
    dt  = ref + timedelta(days=(d - last_day))
    return f'{dt.month}/{dt.day}'

def day_to_date_full(d, last_day, y=2026, m=6, day=11):
    ref = date(y, m, day)
    dt  = ref + timedelta(days=(d - last_day))
    return f'{dt.month}월 {dt.day}일'

def rgb_to_hex(rgb: RGBColor) -> str:
    return f'{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}'

# ── 도형 / 텍스트박스 ──────────────────────────────────────
def add_shape(slide, x, y, w, h, fill=None, line=None, line_w=None):
    shape = slide.shapes.add_shape(1, I(x), I(y), I(w), I(h))
    shape.line.fill.background()
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        if line_w:
            shape.line.width = P(line_w)
    return shape

def add_textbox(slide, x, y, w, h, text, size=9, bold=False, color=None,
                align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = P(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    return txb

# ── 레이아웃 헬퍼 ─────────────────────────────────────────
def add_section_header(slide, x, y, w, h, text):
    add_shape(slide, x, y, 0.05, h, fill=BLUE)
    add_shape(slide, x+0.05, y, w-0.05, h, fill=NAVY)
    add_textbox(slide, x+0.10, y, w-0.10, h, text, size=10, bold=True,
                color=WHITE, align=PP_ALIGN.CENTER)

def add_stat_card(slide, x, y, w, h, label, main, sub, main_color=DARK):
    # 그림자
    add_shape(slide, x+0.025, y+0.025, w, h, fill=RGBColor(0xD1, 0xD9, 0xE8))
    # 카드 배경
    add_shape(slide, x, y, w, h, fill=WHITE,
              line=RGBColor(0xD1, 0xD9, 0xE8), line_w=0.8)
    # 컬러 왼쪽 액센트 바
    add_shape(slide, x, y, 0.05, h, fill=main_color)
    lh = min(0.14, h * 0.26)
    mh = max(0.22, h * 0.52)
    add_textbox(slide, x+0.10, y+0.03,       w-0.14, lh,  label, size=7.5, color=GRAY)
    add_textbox(slide, x+0.10, y+lh+0.03,    w-0.14, mh,  main,  size=15,  bold=True, color=main_color)
    add_textbox(slide, x+0.10, y+h-lh-0.02,  w-0.14, lh,  sub,   size=6.5, color=SLATE)

def add_divider(slide, x, y, w):
    add_shape(slide, x, y, w, 0.01, fill=RGBColor(0xE2, 0xE8, 0xF0))

def add_sub_title(slide, x, y, w, main, sub='', ms=9, ss=8):
    add_textbox(slide, x, y, w*0.48, 0.18, main,
                size=ms, bold=True, color=DARK)
    if sub:
        add_textbox(slide, x+w*0.48, y, w*0.52, 0.18, sub,
                    size=ss, color=GRAY)

# ═══════════════════════════════════════════════════════════
# 차트 XML 유틸 (네이티브 차트 세부 포맷)
# ═══════════════════════════════════════════════════════════

def _xml_solidFill(parent, hex_color: str):
    """parent 아래 solidFill + srgbClr 추가"""
    sf  = etree.SubElement(parent, qn('a:solidFill'))
    clr = etree.SubElement(sf, qn('a:srgbClr'))
    clr.set('val', hex_color.upper().lstrip('#'))
    return sf

def _xml_noFill(parent):
    etree.SubElement(parent, qn('a:noFill'))

def _set_point_colors(chart, hex_colors: list, series_idx=0):
    """막대 차트 시리즈 내 데이터 포인트별 개별 색상 적용 (XML)"""
    ser = chart.series[series_idx]._element
    # cat 요소 앞에 dPt들을 삽입해야 함
    cat = ser.find(qn('c:cat'))
    insert_pos = list(ser).index(cat) if cat is not None else len(list(ser))

    for i, hx in enumerate(hex_colors):
        dPt  = etree.Element(qn('c:dPt'))
        idx  = etree.SubElement(dPt, qn('c:idx'))
        idx.set('val', str(i))
        spPr = etree.SubElement(dPt, qn('c:spPr'))
        _xml_solidFill(spPr, hx)
        ln   = etree.SubElement(spPr, qn('a:ln'))
        _xml_noFill(ln)
        ser.insert(insert_pos + i, dPt)

def _set_line_series(ser_elem, hex_color: str, width_pt=1.5,
                     dash=False, smooth=True, no_marker=True):
    """라인 차트 시리즈 포맷: 색상·두께·대시·스무딩·마커 (XML)"""
    # 기존 spPr 제거 후 새로 삽입
    for old in ser_elem.findall(qn('c:spPr')):
        ser_elem.remove(old)

    spPr = etree.Element(qn('c:spPr'))
    _xml_noFill(spPr)
    ln  = etree.SubElement(spPr, qn('a:ln'))
    ln.set('w', str(int(width_pt * 12700)))  # pt → EMU
    _xml_solidFill(ln, hex_color)
    if dash:
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', 'dash')

    # tx 뒤에 spPr 삽입
    tx = ser_elem.find(qn('c:tx'))
    pos = list(ser_elem).index(tx) + 1 if tx is not None else 2
    ser_elem.insert(pos, spPr)

    # 마커: 기존 것 모두 제거 후 하나만 삽입
    for old_mk in ser_elem.findall(qn('c:marker')):
        ser_elem.remove(old_mk)
    if no_marker:
        mk  = etree.Element(qn('c:marker'))
        sym = etree.SubElement(mk, qn('c:symbol'))
        sym.set('val', 'none')
        ser_elem.insert(pos + 1, mk)

    # 스무딩: 기존 것 모두 제거 후 하나만 추가
    for old_sm in ser_elem.findall(qn('c:smooth')):
        ser_elem.remove(old_sm)
    sm = etree.Element(qn('c:smooth'))
    sm.set('val', '1' if smooth else '0')
    ser_elem.append(sm)

def _set_bar_series_color(chart, series_idx, hex_color: str):
    """막대 시리즈 단색 채우기"""
    ser  = chart.series[series_idx]._element
    spPr = etree.Element(qn('c:spPr'))
    _xml_solidFill(spPr, hex_color)
    ln   = etree.SubElement(spPr, qn('a:ln'))
    _xml_noFill(ln)
    tx   = ser.find(qn('c:tx'))
    pos  = list(ser).index(tx) + 1 if tx is not None else 1
    ser.insert(pos, spPr)

def _invert_cat_axis(chart):
    """카테고리 축 반전 (막대 차트: 상위가 위에 표시되도록)"""
    cat_ax = chart.category_axis._element
    scaling = cat_ax.find(qn('c:scaling'))
    if scaling is None:
        scaling = etree.SubElement(cat_ax, qn('c:scaling'))
        orient  = etree.SubElement(scaling, qn('c:orientation'))
        orient.set('val', 'minMax')
    orient = etree.Element(qn('c:orientation'))
    orient.set('val', 'maxMin')
    scaling.append(orient)

def _fmt_chart(chart, legend=True, legend_pos=XL_LEGEND_POSITION.TOP,
               cat_font=7, val_font=7):
    """공통 차트 포맷"""
    chart.has_title = False          # 차트 제목 비활성화 (시리즈명이 제목으로 표시되는 문제 방지)
    chart.has_legend = legend
    if legend:
        chart.legend.position = legend_pos
        chart.legend.include_in_layout = False
        chart.legend.font.size = P(7)
    chart.category_axis.tick_labels.font.size = P(cat_font)
    chart.value_axis.tick_labels.font.size = P(val_font)

def _remove_gridlines(chart, major=True, minor=True):
    """격자선 제거"""
    ax = chart.value_axis._element
    if major:
        for g in ax.findall(qn('c:majorGridlines')):
            ax.remove(g)
    if minor:
        for g in ax.findall(qn('c:minorGridlines')):
            ax.remove(g)

# ═══════════════════════════════════════════════════════════
# 차트 생성 함수
# ═══════════════════════════════════════════════════════════

def add_trend_chart(slide, trend_data, perf, x, y, w, h):
    """트렌드 라인차트 — 네이티브 LINE"""
    # X축 레이블 과밀 방지: 최대 25개 포인트로 균등 샘플링
    if len(trend_data) > 25:
        n = 25
        idxs = sorted(set(round(i * (len(trend_data) - 1) / (n - 1)) for i in range(n)))
        trend_data = [trend_data[i] for i in idxs]

    days      = [r['day'] for r in trend_data]
    last_day  = perf.get('last_test_day', 75)
    date_lbls = [day_to_date(d, last_day) for d in days]

    def safe(lst): return [v if v is not None else 0 for v in lst]
    true_v  = safe([r.get('true_rate')     for r in trend_data])
    pred_v  = safe([r.get('pred_rate')     for r in trend_data])
    sim_raw = [r.get('sim_pred_rate') for r in trend_data]
    has_sim = any(v is not None and v != 0 for v in sim_raw)
    sim_v   = safe(sim_raw)

    cd = ChartData()
    cd.categories = date_lbls
    cd.add_series('실제 불량률%',    true_v)
    cd.add_series('예측 불량 비율%', pred_v)
    if has_sim:
        cd.add_series('개선 후 비율%', sim_v)

    gf    = slide.shapes.add_chart(XL_CHART_TYPE.LINE, I(x), I(y), I(w), I(h), cd)
    chart = gf.chart

    specs = [
        ('22C55E', 1.5, False, True),
        ('3B82F6', 2.0, False, True),
    ]
    if has_sim:
        specs.append(('F97316', 2.0, True, True))
    for i, (hx, pw, dash, sm) in enumerate(specs):
        _set_line_series(chart.series[i]._element, hx, pw, dash, sm)

    _fmt_chart(chart, legend=True, legend_pos=XL_LEGEND_POSITION.TOP)

    # 카테고리 축 레이블 건너뛰기 (XML 직접 설정 — major_unit은 catAx에서 효과 없음)
    try:
        cat_ax_el = chart.category_axis._element
        skip_el = etree.SubElement(cat_ax_el, qn('c:tickLblSkip'))
        skip_el.set('val', str(max(1, len(days) // 8)))
    except Exception:
        pass

    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = '불량률(%)'
    chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = P(7)

    return gf


def add_shap_chart(slide, shap_data, x, y, w, h):
    """SHAP 가로막대 차트 — 네이티브 BAR_CLUSTERED + 포인트별 색상"""
    if not shap_data:
        return None

    has_shap = shap_data[0].get('mean_abs_shap') is not None
    if has_shap:
        items = sorted(shap_data,
                       key=lambda s: s.get('mean_abs_shap', 0), reverse=True)
        vals  = [s.get('mean_abs_shap', 0) for s in items]
        dirs  = [s.get('high_mean_shap', s.get('effect_norm', 0)) >= 0 for s in items]
        x_lbl = '|SHAP| (mean abs)'
    else:
        items = sorted(shap_data,
                       key=lambda s: abs(s.get('effect_norm', 0)), reverse=True)
        vals  = [abs(s.get('effect_norm', 0)) for s in items]
        dirs  = [s.get('effect_norm', 0) >= 0 for s in items]
        x_lbl = '|effect_norm|'

    items = items[:12]  # 상위 12개만 표시 (너무 많으면 차트가 너무 작아짐)
    feats = [s['feature'] for s in items]
    # 역순: 최상위가 차트 위에 오도록
    feats_r = list(reversed(feats))
    vals_r  = list(reversed(vals))
    dirs_r  = list(reversed(dirs))

    cd = ChartData()
    cd.categories = feats_r
    cd.add_series('', vals_r)

    gf    = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, I(x), I(y), I(w), I(h), cd)
    chart = gf.chart

    # 포인트별 색상: 위험 증가→빨강, 감소→파랑
    hex_colors = ['EF4444' if d else '3B82F6' for d in dirs_r]
    _set_point_colors(chart, hex_colors, series_idx=0)

    _fmt_chart(chart, legend=False)
    _remove_gridlines(chart)

    # x축 레이블 (피처명) 폰트
    chart.category_axis.tick_labels.font.size = P(7)
    chart.value_axis.tick_labels.font.size = P(7)
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = x_lbl
    chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = P(7)

    # 범례 수동 텍스트박스 (빨강=위험증가, 파랑=감소)
    add_shape(slide, x+w-1.6, y+0.03, 0.12, 0.12, fill=RED)
    add_textbox(slide, x+w-1.47, y+0.03, 0.75, 0.14, '위험 증가', size=7, color=SLATE)
    add_shape(slide, x+w-0.72, y+0.03, 0.12, 0.12, fill=BLUE)
    add_textbox(slide, x+w-0.59, y+0.03, 0.55, 0.14, '위험 감소', size=7, color=SLATE)

    return gf


def add_imp_chart(slide, imp_all, fc_names, x, y, w, h):
    """피처 임포턴스 가로막대 — 네이티브 BAR_CLUSTERED + 포인트별 색상"""
    if not imp_all:
        return None

    # 역순 (최상위가 위에)
    items_r = list(reversed(imp_all))
    feats_r = [f['feature'] for f in items_r]
    vals_r  = [f['importance'] for f in items_r]

    cd = ChartData()
    cd.categories = feats_r
    cd.add_series('', vals_r)

    gf    = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, I(x), I(y), I(w), I(h), cd)
    chart = gf.chart

    # SHAP 분석 피처: 인디고, 나머지: 파랑
    hex_colors = ['6366F1' if f in fc_names else '3B82F6' for f in feats_r]
    _set_point_colors(chart, hex_colors, series_idx=0)

    _fmt_chart(chart, legend=False)
    _remove_gridlines(chart)

    chart.category_axis.tick_labels.font.size = P(7)
    chart.value_axis.tick_labels.font.size = P(7)
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = 'Importance (gain)'
    chart.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = P(7)

    # 범례
    add_shape(slide, x+w-1.7, y+0.03, 0.12, 0.12, fill=INDIGO)
    add_textbox(slide, x+w-1.57, y+0.03, 0.95, 0.14, 'SHAP 분석 피처', size=7, color=SLATE)
    add_shape(slide, x+w-0.60, y+0.03, 0.12, 0.12, fill=BLUE)
    add_textbox(slide, x+w-0.47, y+0.03, 0.43, 0.14, '기타', size=7, color=SLATE)

    return gf


def add_hist_chart(slide, sim, x, y, w, h):
    """개선 시뮬레이션 히스토그램 — 네이티브 COLUMN_CLUSTERED"""
    curr_d = sim['current_pred_dist']
    sim_d  = sim['simulated_pred_dist']
    all_v  = curr_d + sim_d
    lo = min(all_v); hi = max(all_v)
    if hi == lo: hi = lo + 0.001

    bins = 8
    import math
    step = (hi - lo) / bins

    curr_cnt = [0] * bins
    sim_cnt  = [0] * bins
    for v in curr_d:
        i = min(int((v - lo) / step), bins - 1)
        curr_cnt[i] += 1
    for v in sim_d:
        i = min(int((v - lo) / step), bins - 1)
        sim_cnt[i] += 1

    lbls = [f'{lo + i*step:.5f}' for i in range(bins)]

    cd = ChartData()
    cd.categories = lbls
    cd.add_series('현재',    curr_cnt)
    cd.add_series('조정 후', sim_cnt)

    gf    = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(x), I(y), I(w), I(h), cd)
    chart = gf.chart

    _set_bar_series_color(chart, 0, 'EF4444')  # 현재: 빨강
    _set_bar_series_color(chart, 1, '22C55E')  # 조정: 초록

    _fmt_chart(chart, legend=True, legend_pos=XL_LEGEND_POSITION.TOP)
    _remove_gridlines(chart)

    chart.category_axis.tick_labels.font.size = P(5.5)
    chart.value_axis.tick_labels.font.size = P(6.5)

    return gf


def add_boxplot_shapes(slide, fc, x, y, w, h):
    """박스플롯 — pptx 도형으로 직접 구현 (고위험 vs 저위험)
    fc[:3] 피처를 나란히 3열 배치
    """
    feats = fc[:3]
    n     = len(feats)
    col_w = w / n
    BOX_H_RATIO = 0.35   # 박스 높이 비율

    for ci, feat in enumerate(feats):
        cx = x + ci * col_w

        # 피처명 + 상관계수
        corr = feat.get('corr_ypred', 0)
        corr_color = RED if abs(corr) > 0.01 else GRAY
        add_textbox(slide, cx + 0.04, y, col_w - 0.08, 0.18,
                    feat['feature'],
                    size=8, bold=True, color=NAVY)
        add_textbox(slide, cx + 0.04, y + 0.17, col_w - 0.08, 0.14,
                    f"r(y_pred)={corr:+.4f}  {feat['agg']} 집계",
                    size=7, color=corr_color)

        # 데이터 범위 계산 (두 box를 함께 스케일)
        hb = feat['high_box']  # [min, q10, q25, q50, q75, q90, max]
        lb = feat['low_box']
        all_pts = hb + lb + [feat.get('normal_lo', min(hb+lb)),
                              feat.get('normal_hi', max(hb+lb))]
        mn = min(all_pts); mx = max(all_pts)
        rng = mx - mn if mx != mn else 1e-6

        chart_x = cx + 0.08
        chart_w = col_w - 0.16
        chart_y = y + 0.34
        chart_h = h - 0.34

        PLOT_W = chart_w * 0.78   # 박스플롯 영역 너비
        PLOT_X = chart_x + chart_w * 0.10   # 왼쪽 여백

        def px(v):   # 값 → x 위치(inch)
            return PLOT_X + (v - mn) / rng * PLOT_W

        def bw(a, b):  # 두 값 사이 너비
            return max(0.005, (b - a) / rng * PLOT_W)

        # 정상 범위 배경
        nlo = feat.get('normal_lo', mn)
        nhi = feat.get('normal_hi', mx)
        add_shape(slide, px(nlo), chart_y,
                  bw(nlo, nhi), chart_h,
                  fill=RGBColor(0xEF, 0xF6, 0xFF))

        row_gap  = chart_h * 0.08
        row_h    = (chart_h - row_gap * 3) / 2
        row_y    = [chart_y + row_gap, chart_y + row_gap * 2 + row_h]
        colors   = [RED, GREEN]
        labels   = ['고위험', '저위험']
        boxes    = [hb, lb]

        for ri, (box, row_top, col, lbl) in enumerate(
                zip(boxes, row_y, colors, labels)):
            q10, q25, q50, q75, q90 = box[1], box[2], box[3], box[4], box[5]
            bx_h = row_h * BOX_H_RATIO
            bx_y = row_top + (row_h - bx_h) / 2

            # 라벨
            add_textbox(slide, chart_x, row_top, chart_w * 0.10, row_h,
                        lbl, size=6.5, bold=True, color=col, align=PP_ALIGN.RIGHT)

            # 왼쪽 수염 (q10 → q25)
            wsk_y = row_top + row_h / 2 - 0.006
            add_shape(slide, px(q10), wsk_y, bw(q10, q25), 0.012, fill=col)

            # IQR 박스 (q25 → q75)
            add_shape(slide, px(q25), bx_y, bw(q25, q75), bx_h,
                      fill=RGBColor(*(int(c * 0.25 + 0xFF * 0.75)
                                      for c in [col[0], col[1], col[2]])),
                      line=col, line_w=0.6)

            # 중앙값 선
            add_shape(slide, px(q50) - 0.010, bx_y, 0.020, bx_h, fill=WHITE)

            # 오른쪽 수염 (q75 → q90)
            add_shape(slide, px(q75), wsk_y, bw(q75, q90), 0.012, fill=col)

            # 중앙값 텍스트
            add_textbox(slide, px(q50) + 0.02, row_top, 0.55, row_h,
                        f'Med:{q50:.4f}', size=6, color=SLATE)

        # 정상 범위 텍스트
        add_textbox(slide, chart_x, chart_y + chart_h - 0.16, chart_w, 0.16,
                    f"정상 범위: {feat['normal_lo']} ~ {feat['normal_hi']}",
                    size=6.5, color=SLATE)

        # 열 구분선 (마지막 열 제외)
        if ci < n - 1:
            add_shape(slide, cx + col_w - 0.01, y, 0.01, h,
                      fill=RGBColor(0xE2, 0xE8, 0xF0))

# ═══════════════════════════════════════════════════════════
# 보고서 요소
# ═══════════════════════════════════════════════════════════

def add_rmse_display(slide, x, y, w, h, val_rmse):
    add_shape(slide, x+0.025, y+0.025, w, h, fill=RGBColor(0xD1, 0xD9, 0xE8))
    add_shape(slide, x, y, w, h, fill=WHITE,
              line=RGBColor(0xD1, 0xD9, 0xE8), line_w=0.8)
    color = GREEN if val_rmse < 0.006 else AMBER
    add_shape(slide, x, y, 0.05, h, fill=color)
    num_h = max(0.20, h - 0.18)
    add_textbox(slide, x+0.12, y+0.03, w-0.16, num_h,
                f'{val_rmse:.6f}', size=20, bold=True, color=color)
    add_textbox(slide, x+0.12, y+num_h+0.02, w-0.16, 0.14,
                'Validation RMSE', size=7.5, color=GRAY)


def add_wafer_table(slide, x, y, w, h, wafer_risk):
    high = [r for r in wafer_risk if r['pred_rate'] >= 30]
    if not high:
        high = sorted(wafer_risk, key=lambda r: r['pred_rate'], reverse=True)[:6]
    high = sorted(high, key=lambda r: r['id'])
    rows_data = high[:6]
    n_rows    = len(rows_data) + 1

    tbl = slide.shapes.add_table(n_rows, 3, I(x), I(y), I(w), I(h)).table
    tbl.columns[0].width = I(w * 0.42)
    tbl.columns[1].width = I(w * 0.38)
    tbl.columns[2].width = I(w * 0.20)

    def set_cell(r, c, text, bold=False, color=DARK, bg=None,
                 align=PP_ALIGN.LEFT, size=8.5):
        cell = tbl.cell(r, c)
        cell.text = text
        para = cell.text_frame.paragraphs[0]
        para.alignment = align
        run = para.runs[0] if para.runs else para.add_run()
        run.font.size = P(size)
        run.font.bold = bold
        if color: run.font.color.rgb = color
        if bg:
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            sf   = etree.SubElement(tcPr, qn('a:solidFill'))
            clr  = etree.SubElement(sf,   qn('a:srgbClr'))
            clr.set('val', f'{bg[0]:02X}{bg[1]:02X}{bg[2]:02X}')

    for c, h_txt in enumerate(['Lot-Wafer', '예측 불량률', '유닛 수']):
        set_cell(0, c, h_txt, bold=True, color=SLATE,
                 bg=(0xF8, 0xFA, 0xFC), align=PP_ALIGN.CENTER)

    for ri, row in enumerate(rows_data, 1):
        rate = row['pred_rate']
        col  = (0xDC,0x26,0x26) if rate>=60 else \
               (0xF5,0x9E,0x0B) if rate>=40 else (0x3B,0x82,0xF6)
        set_cell(ri, 0, row['id'],           bold=True,  color=RGBColor(0x1E,0x29,0x3B))
        set_cell(ri, 1, f'{rate}%',          bold=True,  color=RGBColor(*col), align=PP_ALIGN.CENTER)
        set_cell(ri, 2, str(row.get('n','')),color=GRAY, align=PP_ALIGN.CENTER)


def add_improvement_table(slide, x, y, w, h, feat0):
    sim = feat0['sim']
    rows_data = [
        ('고위험 중앙값',       str(feat0['high_median'])),
        ('목표 (저위험 중앙값)', str(feat0['normal_median'])),
        ('정상 범위',           f"{feat0['normal_lo']} ~ {feat0['normal_hi']}"),
        ('조정량',              f"{'+' if sim['delta_feat']>=0 else ''}{sim['delta_feat']}"),
    ]
    tbl = slide.shapes.add_table(
        len(rows_data), 2,
        I(x), I(y), I(w * 0.56), I(h)
    ).table
    tbl.columns[0].width = I(w * 0.32)
    tbl.columns[1].width = I(w * 0.24)

    def sc(r, c, text, bold=False, color=DARK, size=8):
        cell = tbl.cell(r, c)
        cell.text = ''
        para = cell.text_frame.paragraphs[0]
        run = para.add_run()
        run.text = text
        run.font.size = P(size)
        run.font.bold = bold
        if color: run.font.color.rgb = color

    for ri, (k, v) in enumerate(rows_data):
        sc(ri, 0, k, color=SLATE)
        sc(ri, 1, v, bold=True)

    # 오른쪽 변화 박스
    bx = x + w * 0.58
    improved = sim['new_pred_est'] < sim['current_pred_median']
    add_shape(slide, bx, y, w*0.42, h,
              fill=RGBColor(0xF0,0xFD,0xF4),
              line=RGBColor(0xBB,0xF7,0xD0), line_w=0.5)
    add_textbox(slide, bx+0.05, y+0.02, w*0.40, 0.14,
                '예상 y_pred 변화', size=7, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, bx+0.05, y+0.14, w*0.13, 0.14,
                '현재', size=6.5, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, bx+0.05, y+0.25, w*0.13, 0.24,
                str(sim['current_pred_median']),
                size=11, bold=True, color=RED, align=PP_ALIGN.CENTER)
    add_textbox(slide, bx+w*0.19, y+0.25, w*0.06, 0.24,
                '->', size=11, color=SLATE, align=PP_ALIGN.CENTER)
    adj_c = GREEN if improved else RED
    add_textbox(slide, bx+w*0.25, y+0.14, w*0.15, 0.14,
                '조정 후', size=6.5, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide, bx+w*0.25, y+0.25, w*0.15, 0.24,
                str(sim['new_pred_est']),
                size=11, bold=True, color=adj_c, align=PP_ALIGN.CENTER)
    dsgn = '+' if sim['delta_pred'] >= 0 else ''
    add_textbox(slide, bx+0.05, y+0.50, w*0.40, 0.14,
                f"delta {dsgn}{sim['delta_pred']} {'개선' if improved else '변화없음'}",
                size=7, color=GREEN if improved else SLATE, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════
# 커스텀 섹션 슬라이드 (HTML과 동일하게 한 슬라이드에 그리드 배치)
# ═══════════════════════════════════════════════════════════

_CS_PALETTE = ['3B82F6', 'EF4444', '22C55E', 'F59E0B', '8B5CF6',
               '06B6D4', 'F97316', 'EC4899', '64748B', '84CC16']


def _render_cs_chart(slide, sec, x, y, w, h):
    """커스텀 섹션 차트를 지정 영역에 그린다 (table 제외)."""
    chart_type = sec.get('chart_type', 'bar')
    labels     = sec.get('labels', [])
    datasets   = sec.get('datasets', [])
    if not datasets:
        add_textbox(slide, x, y, w, h, '(데이터 없음)', size=8, color=GRAY)
        return

    pptx_type_map = {
        'line':     XL_CHART_TYPE.LINE,
        'pie':      XL_CHART_TYPE.PIE,
        'doughnut': XL_CHART_TYPE.DOUGHNUT,
        'scatter':  XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
    }
    if chart_type == 'bar' and sec.get('horizontal'):
        pptx_type = XL_CHART_TYPE.BAR_CLUSTERED
    elif chart_type == 'bar':
        pptx_type = XL_CHART_TYPE.COLUMN_CLUSTERED
    else:
        pptx_type = pptx_type_map.get(chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)

    cd = ChartData()
    cd.categories = labels
    for ds in datasets:
        cd.add_series(ds.get('label', ''), ds.get('data', []))

    gf    = slide.shapes.add_chart(pptx_type, I(x), I(y), I(w), I(h), cd)
    chart = gf.chart
    chart.has_title = False

    has_labels = any(ds.get('label') for ds in datasets)
    chart.has_legend = len(datasets) > 1 or has_labels
    if chart.has_legend:
        chart.legend.position          = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size         = P(7)

    for i, ds in enumerate(datasets):
        raw_color = ds.get('color', '#' + _CS_PALETTE[i % len(_CS_PALETTE)])
        hex_color = raw_color.lstrip('#')
        if chart_type in ('pie', 'doughnut'):
            colors = [_CS_PALETTE[j % len(_CS_PALETTE)] for j in range(len(labels))]
            _set_point_colors(chart, colors, series_idx=i)
        elif chart_type == 'line':
            _set_line_series(chart.series[i]._element, hex_color, 2.0)
        else:
            _set_bar_series_color(chart, i, hex_color)

    if chart_type not in ('pie', 'doughnut'):
        try:
            chart.category_axis.tick_labels.font.size = P(7)
            chart.value_axis.tick_labels.font.size    = P(7)
            _remove_gridlines(chart)
        except Exception:
            pass


def _render_cs_table(slide, sec, x, y, w, h):
    """커스텀 섹션 테이블을 지정 영역에 그린다."""
    columns  = sec.get('columns', [])
    rows     = sec.get('rows', [])
    if not columns or not rows:
        add_textbox(slide, x, y, w, 0.4, '(테이블 데이터 없음)', size=8, color=GRAY)
        return

    max_rows = min(len(rows), 20)
    n_rows   = max_rows + 1
    n_cols   = len(columns)

    tbl = slide.shapes.add_table(n_rows, n_cols, I(x), I(y), I(w), I(h)).table
    tbl.rows[0].height = I(0.26)
    for ci, col in enumerate(columns):
        cell = tbl.cell(0, ci)
        cell.text = str(col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p   = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.runs[0] if p.runs else p.add_run()
        run.font.size      = P(7)
        run.font.bold      = True
        run.font.color.rgb = WHITE

    for ri, row in enumerate(rows[:max_rows]):
        row_data = row if isinstance(row, list) else [row.get(c, '') for c in columns]
        tbl.rows[ri + 1].height = I(0.22)
        bg = LTGRAY if ri % 2 == 0 else WHITE
        for ci in range(n_cols):
            cell = tbl.cell(ri + 1, ci)
            cell.text = str(row_data[ci]) if ci < len(row_data) else ''
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            p   = cell.text_frame.paragraphs[0]
            run = p.runs[0] if p.runs else p.add_run()
            run.font.size      = P(7)
            run.font.color.rgb = DARK if ci == 0 else SLATE


def _render_cs_card(slide, sec, x, y, w, h):
    """커스텀 섹션 카드 한 개를 (x,y,w,h) 영역에 그린다 — HTML 카드와 동일 구조."""
    SEC_HDR_H = 0.26
    PAD       = 0.07

    # 카드 배경 (흰색, 회색 테두리)
    add_shape(slide, x, y, w, h, fill=WHITE,
              line=RGBColor(0xCB, 0xD5, 0xE1), line_w=0.6)

    # 카드 헤더 (왼쪽 파랑 강조 + 네이비 배경)
    add_shape(slide, x,        y, 0.05,       SEC_HDR_H, fill=BLUE)
    add_shape(slide, x + 0.05, y, w - 0.05,  SEC_HDR_H, fill=NAVY)
    title = sec.get('title', '추가 분석')
    add_textbox(slide, x + 0.10, y + 0.04, w - 0.14, SEC_HDR_H - 0.06,
                title, size=8, bold=True, color=WHITE)

    # 설명 텍스트 (있을 때)
    desc  = sec.get('description', '')
    cy    = y + SEC_HDR_H + PAD
    if desc:
        add_textbox(slide, x + PAD, cy, w - PAD * 2, 0.18,
                    desc, size=7, color=SLATE)
        cy += 0.20

    # 차트 / 테이블 영역
    content_h = max(0.4, y + h - cy - PAD)
    chart_type = sec.get('chart_type', 'bar')
    if chart_type == 'table':
        _render_cs_table(slide, sec, x + PAD, cy, w - PAD * 2, content_h)
    else:
        _render_cs_chart(slide, sec, x + PAD, cy, w - PAD * 2, content_h)


def add_all_custom_sections_slide(prs: 'Presentation', custom_sections: list):
    """모든 custom_sections를 HTML과 동일하게 한 슬라이드에 그리드로 배치."""
    if not custom_sections:
        return

    from datetime import datetime

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    today = datetime.now().strftime('%Y년 %m월 %d일')

    # 배경
    add_shape(slide, 0, 0, 13.33, 7.5, fill=RGBColor(0xF1, 0xF5, 0xF9))

    # 헤더
    HDR_H = 0.56
    add_shape(slide, 0, 0, 13.33, HDR_H, fill=NAVY)
    add_shape(slide, 0, 0, 0.10,  HDR_H, fill=BLUE)
    add_textbox(slide, 0.21, 0.05, 12.90, 0.32,
                '추가 분석 섹션', size=19, bold=True, color=WHITE)
    add_textbox(slide, 0.21, 0.36, 12.90, 0.16, today,
                size=7.5, color=RGBColor(0xBA, 0xCC, 0xF0))

    # 푸터
    F_Y = 7.5 - 0.22
    add_shape(slide, 0, F_Y - 0.02, 13.33, 0.025, fill=NAVY)
    add_textbox(slide, 0.15, F_Y + 0.01, 5.0,  0.20, today, size=7, color=SLATE)
    add_textbox(slide, 5.15, F_Y + 0.01, 3.5,  0.20,
                'Field Health Prediction Model v1.0',
                size=7, color=SLATE, align=PP_ALIGN.CENTER)
    add_textbox(slide, 8.65, F_Y + 0.01, 4.53, 0.20,
                'We Do Technology  |  SK hynix',
                size=7, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)

    # 그리드 레이아웃 계산
    LM      = 0.15
    GAP     = 0.14
    BODY_Y  = HDR_H + 0.14
    AVAIL_H = F_Y - 0.04 - BODY_Y
    INNER_W = 13.33 - LM * 2

    n    = len(custom_sections)
    cols = 1 if n == 1 else 2
    rows = (n + cols - 1) // cols      # ceil(n / cols)

    col_w = (INNER_W - GAP * (cols - 1)) / cols
    row_h = (AVAIL_H - GAP * (rows - 1)) / rows

    for idx, sec in enumerate(custom_sections):
        col = idx % cols
        row = idx // cols
        cx  = LM + col * (col_w + GAP)
        cy  = BODY_Y + row * (row_h + GAP)
        try:
            _render_cs_card(slide, sec, cx, cy, col_w, row_h)
        except Exception as e:
            sys.stdout.buffer.write(
                f'[경고] custom_section[{idx}] 렌더링 실패: {e}\n'.encode('utf-8'))

    return slide


# ═══════════════════════════════════════════════════════════
# 메인 PPT 빌드
# ═══════════════════════════════════════════════════════════

def build_ppt(data: dict, out_path: str):
    from datetime import datetime
    p          = data['perf']
    fc         = data['feature_comparison']
    shap_data  = data.get('shap_data', [])
    wafer_risk = data.get('wafer_risk', [])
    trend_data = data.get('trend_data', [])
    feat0      = fc[0]
    today      = datetime.now().strftime('%Y년 %m월 %d일')
    last_day   = p.get('last_test_day', 75)
    thr        = p.get('pred_defect_thr', p.get('high_thr', 0.003))

    prs = Presentation()
    prs.slide_width  = SW
    prs.slide_height = SH
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # ── 전체 배경 ──────────────────────────────────────────
    add_shape(slide, 0, 0, 13.33, 7.5, fill=RGBColor(0xF1, 0xF5, 0xF9))

    # ── 레이아웃 상수 ──────────────────────────────────────
    LM, RM   = 0.15, 0.15
    INNER_W  = 13.33 - LM - RM
    F_Y      = 7.5 - 0.22
    GAP      = 0.14
    L_W      = 6.10
    R_W      = INNER_W - L_W - GAP
    L_X      = LM
    R_X      = LM + L_W + GAP

    # ── 네이비 헤더 배너 ───────────────────────────────────
    HDR_H = 0.56
    add_shape(slide, 0, 0, 13.33, HDR_H, fill=NAVY)
    add_shape(slide, 0, 0, 0.10, HDR_H, fill=BLUE)   # 왼쪽 파랑 강조 바
    add_textbox(slide, LM+0.06, 0.05, INNER_W*0.74, 0.32,
                'Field Health 불량 예측 분석 보고서',
                size=19, bold=True, color=WHITE)
    add_textbox(slide, LM+0.06, 0.36, INNER_W*0.74, 0.16,
                f'{today}  ·  Two-Stage Model  ·  Val RMSE {p["val_rmse"]:.6f}  ·  임계값 y_pred ≥ {thr}',
                size=7.5, color=RGBColor(0xBA, 0xCC, 0xF0))
    add_shape(slide, 13.33-RM-1.12, 0.14, 1.12, 0.28,
              fill=RGBColor(0xDC, 0x26, 0x26),
              line=RGBColor(0x99, 0x1B, 0x1B), line_w=0.8)
    add_textbox(slide, 13.33-RM-1.12, 0.14, 1.12, 0.28,
                '대외비', size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # ── 본문 ──────────────────────────────────────────────
    BODY_Y = HDR_H + 0.12

    # Custom sections: 메인+CS 전체를 비례 스케일링해 한 슬라이드에 배치
    custom_sections = data.get('custom_sections', [])
    n_cs    = len(custom_sections)
    cs_cols = 2 if n_cs > 1 else 1
    cs_rows = (n_cs + cs_cols - 1) // cs_cols if n_cs > 0 else 0

    AVAIL_FULL = F_Y - BODY_Y                      # 6.60" — CS 없을 때 풀 높이
    if n_cs > 0:
        # 이상적 높이: 메인 6.60" + CS 행(1.80" + 0.12" 갭) × rows
        total_ideal = AVAIL_FULL + cs_rows * 1.92
        _s       = AVAIL_FULL / total_ideal         # 압축 비율
        AVAIL    = AVAIL_FULL * _s
        CS_ROW_H = 1.80 * _s
        CS_GAP   = 0.12 * _s
    else:
        AVAIL    = AVAIL_FULL
        CS_ROW_H = 0.0
        CS_GAP   = 0.0
    _scale   = AVAIL / AVAIL_FULL                  # 1.0 if no CS
    MAIN_BOT = BODY_Y + AVAIL

    # 섹션 배경 패널 (먼저 그려야 위에 내용이 얹힘)
    add_shape(slide, L_X, BODY_Y, L_W, AVAIL,
              fill=WHITE, line=RGBColor(0xCB, 0xD5, 0xE1), line_w=0.6)
    add_shape(slide, R_X, BODY_Y, R_W, AVAIL,
              fill=WHITE, line=RGBColor(0xCB, 0xD5, 0xE1), line_w=0.6)

    # ═══════════ 왼쪽: [불량 현황] ════════════════════════
    cy = BODY_Y

    add_section_header(slide, L_X, cy, L_W, 0.26, '[ 불량 현황 ]')
    cy += 0.28

    # 통계 카드 3개
    n_unit   = p.get('n_val', 0)
    n_high   = p.get('n_high', 0)
    pred_pct = f'{n_high/n_unit*100:.1f}%' if n_unit else '0.0%'
    cw = (L_W - 0.12) / 3
    CARD_H = 0.56
    add_stat_card(slide, L_X+0.04,       cy, cw-0.04, CARD_H,
                  '기준일', day_to_date_full(last_day, last_day), '최신 1일 기준', NAVY)
    add_stat_card(slide, L_X+0.04+cw,    cy, cw-0.04, CARD_H,
                  '분석 유닛', f'{n_unit}개', 'last lot', DARK)
    pred_col = RED if n_high > 0 else GREEN
    add_stat_card(slide, L_X+0.04+cw*2,  cy, cw-0.04, CARD_H,
                  '예측 불량', f'{n_high}개',
                  f'{pred_pct} (≥{thr})', pred_col)
    cy += CARD_H + 0.10

    # 모델 성능 (RMSE)
    add_sub_title(slide, L_X+0.04, cy, L_W-0.08,
                  '1. 모델 성능',
                  f'Study: {p.get("study_id", "")}',
                  ms=8, ss=7)
    cy += 0.18
    RMSE_H = 0.46
    add_rmse_display(slide, L_X+0.04, cy, L_W-0.08, RMSE_H, p['val_rmse'])
    cy += RMSE_H + 0.10
    add_divider(slide, L_X+0.04, cy, L_W-0.08); cy += 0.10

    # 트렌드 차트
    add_sub_title(slide, L_X+0.04, cy, L_W-0.08,
                  f'2. 불량률 트렌드  (Train {p["train_days"]}일 / Val {p["val_days"]}일 / Test {p["test_days"]}일)',
                  '주황 점선: 개선 후', ms=8, ss=7)
    cy += 0.18

    TREND_H = max(0.80, 1.32 * _scale)
    if trend_data:
        add_trend_chart(slide, trend_data, p, L_X+0.04, cy, L_W-0.08, TREND_H)
    cy += TREND_H + 0.08

    add_divider(slide, L_X+0.04, cy, L_W-0.08); cy += 0.10

    # SHAP 차트
    shap_meta = data.get('shap_meta', {})
    r2_txt = f' · R²={shap_meta["surrogate_r2"]}' if shap_meta else ''
    add_sub_title(slide, L_X+0.04, cy, L_W-0.08,
                  '3. SHAP 분석',
                  f'TreeSHAP · 서로게이트 LightGBM{r2_txt} · TOP {min(len(shap_data), 12)}',
                  ms=8, ss=7)
    cy += 0.18

    # 남은 왼쪽 공간을 SHAP 차트로 가득 채움
    SHAP_H = max(AVAIL * 0.17, (BODY_Y + AVAIL) - cy - 0.06)
    if shap_data:
        add_shap_chart(slide, shap_data, L_X+0.04, cy, L_W-0.08, SHAP_H)

    # ═══════════ 오른쪽: [이력 및 대책] ═══════════════════
    cy = BODY_Y

    add_section_header(slide, R_X, cy, R_W, 0.26, '[ 이력 및 대책 ]')
    cy += 0.28

    # 1. Wafer 위험 테이블
    add_sub_title(slide, R_X+0.04, cy, R_W-0.08,
                  '1. 위험 집중 Wafer',
                  'Test 기준 · 예측 불량률 ≥ 30%',
                  ms=8, ss=7)
    cy += 0.18

    if wafer_risk:
        n_wr  = min(6, len(wafer_risk))
        ROW_H = 0.170
        tbl_h = 0.24 + n_wr * ROW_H
        add_wafer_table(slide, R_X+0.04, cy, R_W-0.08, tbl_h, wafer_risk)
        cy += tbl_h + 0.08
    else:
        add_textbox(slide, R_X+0.04, cy, R_W-0.08, 0.16, '데이터 없음', size=7.5, color=GRAY)
        cy += 0.24

    add_divider(slide, R_X+0.04, cy, R_W-0.08); cy += 0.10

    # 2. 고위험 피처 분포 (박스플롯)
    add_sub_title(slide, R_X+0.04, cy, R_W-0.08,
                  '2. 고위험 피처 분포',
                  f'고위험(≥{p.get("high_thr", "")}) vs 저위험(<{p.get("low_thr", "")})',
                  ms=8, ss=7)
    cy += 0.18

    # 박스플롯 + 개선방안의 총 가용 높이를 계산해 배분 (6:4)
    right_remain = (BODY_Y + AVAIL) - cy - 0.18 - 0.10 - 0.18 - 0.06  # 섹션 제목·구분선 제외
    BP_H  = max(AVAIL * 0.21, right_remain * 0.58)
    IMP_H = max(AVAIL * 0.08, right_remain * 0.42)

    if fc:
        add_boxplot_shapes(slide, fc, R_X+0.04, cy, R_W-0.08, BP_H)
        cy += BP_H + 0.08

    add_divider(slide, R_X+0.04, cy, R_W-0.08); cy += 0.10

    # 3. 개선 방안
    add_sub_title(slide, R_X+0.04, cy, R_W-0.08,
                  '3. 개선 방안',
                  f'{feat0["feature"]}  ({feat0["agg"]} 집계)',
                  ms=8, ss=7.5)
    cy += 0.18
    add_improvement_table(slide, R_X+0.04, cy, R_W-0.08, IMP_H, feat0)

    # ── 푸터 ──────────────────────────────────────────────
    add_shape(slide, 0, F_Y-0.02, 13.33, 0.025, fill=NAVY)
    add_textbox(slide, LM, F_Y+0.01, INNER_W*0.38, 0.20,
                f'{today}  ·  Two-Stage Model  ·  임계값 y_pred ≥ {thr}',
                size=7, color=SLATE)
    add_textbox(slide, LM+INNER_W*0.28, F_Y+0.01, INNER_W*0.44, 0.20,
                'Field Health Prediction Model v1.0  |  WT Quality Analysis  |  Page 1 of 1',
                size=7, color=SLATE, align=PP_ALIGN.CENTER)
    add_textbox(slide, LM+INNER_W*0.72, F_Y+0.01, INNER_W*0.28, 0.20,
                'We Do Technology  |  SK hynix',
                size=7, bold=True, color=NAVY, align=PP_ALIGN.RIGHT)

    # Custom sections — 메인 슬라이드 하단에 직접 배치 (새 슬라이드 아님)
    if custom_sections:
        CS_Y     = MAIN_BOT + CS_GAP * 0.5
        cs_col_w = (INNER_W - CS_GAP * (cs_cols - 1)) / cs_cols
        for idx, sec in enumerate(custom_sections):
            col = idx % cs_cols
            row = idx // cs_cols
            cx  = LM + col * (cs_col_w + CS_GAP)
            cy  = CS_Y + row * (CS_ROW_H + CS_GAP)
            try:
                _render_cs_card(slide, sec, cx, cy, cs_col_w, CS_ROW_H)
            except Exception as e:
                sys.stdout.buffer.write(
                    f'[경고] custom_section[{idx}] 렌더링 실패: {e}\n'.encode('utf-8'))

    prs.save(out_path)
    sys.stdout.buffer.write(f'[완료] {out_path}\n'.encode('utf-8'))


# ── 진입점 ────────────────────────────────────────────────
if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--data', default='report_data.json')
    parser.add_argument('--out',  default='report_v3.pptx')
    args = parser.parse_args()

    base      = os.path.dirname(os.path.abspath(__file__))
    data_path = args.data if os.path.isabs(args.data) else os.path.join(base, args.data)
    out_path  = args.out  if os.path.isabs(args.out)  else os.path.join(base, args.out)

    with open(data_path, encoding='utf-8') as f:
        data = json.load(f)

    build_ppt(data, out_path)
